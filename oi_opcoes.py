# oi_opcoes.py — Milho Trader v1.0 (Fase 3B)
# Barreiras de opções CCM: Call Wall, Put Wall, Max Pain — a partir da grade
# de opções exportada do Profit/Genial (arquivo CCM_OP.xlsx).
#
# IMPORTANTE — por que este módulo NÃO usa as colunas CONTRATO/ANO/TIPO/Strike
# do arquivo exportado: elas vêm de fórmulas do Excel que, na amostra real
# (11/07/2026), estavam incorretas em ~53% das linhas (ex.: um instrumento com
# ticker "CCMU26C007400" — vencimento U/26 pelo próprio código — aparecia com
# CONTRATO='X' na planilha). O ticker em "Instrumento financeiro" é a única
# fonte confiável: segue o padrão CCM + mês(1 letra) + ano(2 dígitos) +
# tipo(C/P) + strike×100 (6 dígitos), ex.: CCMU26C006700 = CCM, Set/2026,
# Call, strike R$67,00. Todos os campos usados aqui são derivados desse
# ticker via regex, nunca das colunas auxiliares.

import os
import re
import glob
import time

import openpyxl

from leitor_csv import ler_arquivo, ultimo_valor

_PAT_TICKER = re.compile(r'^CCM([A-Z])(\d{2})([CP])(\d+)$')

_MES_ORDEM = {
    'F': 1, 'G': 2, 'H': 3, 'J': 4, 'K': 5, 'M': 6,
    'N': 7, 'Q': 8, 'U': 9, 'V': 10, 'X': 11, 'Z': 12,
}

OI_VAZIO = {
    'call_wall':            None,
    'put_wall':             None,
    'max_pain':             None,
    'call_wall_agregado':   None,
    'put_wall_agregado':    None,
    'grade_agregada':       [],
    'expiracoes_agregadas': [],
    'fonte':                'manual',
    'vencimento_opcoes':    None,
    'total_oi_calls':       0,
    'total_oi_puts':        0,
    'arquivo':              None,
    'grade':                [],
}


def _encontrar_arquivo_op(pasta: str) -> str:
    """
    Localiza a grade de opções mais recente na pasta. Exclui arquivos de
    lock do Excel (prefixo '~$'), que aparecem quando o arquivo está aberto
    no Excel no momento da exportação/leitura.
    """
    padrao = os.path.join(pasta, 'CCM_OP*.xlsx')
    encontrados = [f for f in glob.glob(padrao) if not os.path.basename(f).startswith('~$')]
    if not encontrados:
        return None
    return sorted(encontrados)[-1]


def _extrair_grade(path: str) -> list:
    """
    Abre o .xlsx e varre TODAS as abas (o arquivo real tem 3: dados brutos,
    dados com colunas auxiliares, e um pivô-resumo — a ordem/nome pode variar
    entre exportações). Usa a aba que produzir mais linhas reconhecíveis pelo
    padrão de ticker CCM, e ignora as demais.

    Retry curto para o caso comum de o arquivo estar aberto no Excel no
    instante exato da leitura (mesmo padrão defensivo dos outros leitores
    do projeto).
    """
    ultima_excecao = None
    wb = None
    for tentativa in range(3):
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            break
        except Exception as e:
            ultima_excecao = e
            if tentativa < 2:
                time.sleep(0.5)
    if wb is None:
        raise ultima_excecao

    melhor = []
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            continue

        header = [str(c).strip().lower() if c is not None else '' for c in linhas[0]]
        idx_inst = next((i for i, h in enumerate(header) if 'instrumento' in h), None)
        idx_oi   = next((i for i, h in enumerate(header) if 'contratos em aberto' in h and 'varia' not in h), None)

        if idx_inst is not None and idx_oi is not None:
            corpo = linhas[1:]
        else:
            idx_inst, idx_oi = 0, 1
            corpo = linhas

        registros = []
        for r in corpo:
            if not r or len(r) <= max(idx_inst, idx_oi):
                continue
            ticker, oi = r[idx_inst], r[idx_oi]
            if ticker is None or oi is None:
                continue
            m = _PAT_TICKER.match(str(ticker).strip())
            if not m:
                continue
            mes, ano, tipo, strike_raw = m.groups()
            try:
                oi_val = float(oi)
            except (TypeError, ValueError):
                continue
            registros.append({
                'ticker': str(ticker).strip(),
                'mes':    mes,
                'ano':    ano,
                'tipo':   tipo,
                'strike': int(strike_raw) / 100.0,
                'oi':     oi_val,
            })

        if len(registros) > len(melhor):
            melhor = registros

    return melhor


def _chave_mes_ano(mes: str, ano: str) -> tuple:
    return (2000 + int(ano), _MES_ORDEM.get(mes, 99))


def _extrair_mes_ano_vencimento(codigo_vencimento: str) -> tuple:
    """Ex.: 'CCMU26' → ('U', '26'). Códigos inesperados caem no except do chamador."""
    return codigo_vencimento[3], codigo_vencimento[4:6]


def calcular_max_pain(calls: dict, puts: dict):
    """
    Strike que minimiza o payout total agregado a titulares de calls+puts
    caso o ativo feche exatamente nesse preço no vencimento. Definição
    padrão de mercado, baseada em OI bruto (não gamma-weighted).
    """
    strikes = sorted(set(list(calls.keys()) + list(puts.keys())))
    if not strikes:
        return None
    melhor_strike, melhor_pain = None, None
    for s in strikes:
        pain = (sum(oi * (s - k) for k, oi in calls.items() if k < s) +
                sum(oi * (k - s) for k, oi in puts.items() if k > s))
        if melhor_pain is None or pain < melhor_pain:
            melhor_pain, melhor_strike = pain, s
    return melhor_strike


def calcular_wall(dic_oi: dict, preco_atual: float, lado: str):
    """
    Call Wall = strike de CALL com maior OI ACIMA do preço atual (resistência).
    Put Wall  = strike de PUT com maior OI ABAIXO do preço atual (suporte).
    Se não houver strikes do lado esperado (raro — preço fora de toda a
    grade), cai para o maior OI da grade inteira em vez de retornar vazio.
    """
    if not dic_oi:
        return None
    if lado == 'call':
        candidatos = {k: v for k, v in dic_oi.items() if k >= preco_atual}
    else:
        candidatos = {k: v for k, v in dic_oi.items() if k <= preco_atual}
    base = candidatos if candidatos else dic_oi
    return max(base, key=base.get)


def analisar_oi_opcoes(pasta: str, vencimento_ativo: str, preco_atual: float) -> dict:
    """
    Ponto de entrada. Localiza o arquivo, extrai a grade, filtra pelo
    vencimento ativo do futuro (mesmo mês/ano do CCM em operação) e calcula
    Call Wall, Put Wall e Max Pain. Se o vencimento exato não tiver opções
    na grade, usa o vencimento disponível mais próximo cronologicamente.
    """
    path = _encontrar_arquivo_op(pasta)
    if path is None:
        return dict(OI_VAZIO)

    registros = _extrair_grade(path)
    if not registros:
        saida = dict(OI_VAZIO)
        saida['arquivo'] = os.path.basename(path)
        return saida

    mes_alvo, ano_alvo = _extrair_mes_ano_vencimento(vencimento_ativo)
    alvo = [r for r in registros if r['mes'] == mes_alvo and r['ano'] == ano_alvo]

    if not alvo:
        disponiveis = sorted(set((r['mes'], r['ano']) for r in registros), key=lambda t: _chave_mes_ano(*t))
        alvo_idx = _chave_mes_ano(mes_alvo, ano_alvo)
        alvo_ord = alvo_idx[0] * 12 + alvo_idx[1]
        mes_alvo, ano_alvo = min(
            disponiveis,
            key=lambda t: abs((_chave_mes_ano(*t)[0] * 12 + _chave_mes_ano(*t)[1]) - alvo_ord)
        )
        alvo = [r for r in registros if r['mes'] == mes_alvo and r['ano'] == ano_alvo]

    calls = {r['strike']: r['oi'] for r in alvo if r['tipo'] == 'C'}
    puts  = {r['strike']: r['oi'] for r in alvo if r['tipo'] == 'P'}

    todos_strikes = sorted(set(list(calls.keys()) + list(puts.keys())))
    grade = [
        {'strike': s, 'oi_call': calls.get(s, 0), 'oi_put': puts.get(s, 0)}
        for s in todos_strikes
    ]

    # Agregado somando TODAS as expirações presentes no arquivo (não só o
    # vencimento ativo) — mostra os níveis que o mercado inteiro está
    # defendendo, não só quem opera aquele mês específico. Decidido com o
    # usuário em 11/07/2026 como a visão principal para destacar no
    # gráfico de candles (Seção "Opções" do dashboard).
    calls_agg, puts_agg = {}, {}
    for r in registros:
        alvo_dict = calls_agg if r['tipo'] == 'C' else puts_agg
        alvo_dict[r['strike']] = alvo_dict.get(r['strike'], 0) + r['oi']
    strikes_agg = sorted(set(list(calls_agg.keys()) + list(puts_agg.keys())))
    grade_agregada = [
        {'strike': s, 'oi_call': calls_agg.get(s, 0), 'oi_put': puts_agg.get(s, 0)}
        for s in strikes_agg
    ]

    return {
        'call_wall':         calcular_wall(calls, preco_atual, 'call'),
        'put_wall':          calcular_wall(puts, preco_atual, 'put'),
        'max_pain':          calcular_max_pain(calls, puts),
        'call_wall_agregado': calcular_wall(calls_agg, preco_atual, 'call'),
        'put_wall_agregado':  calcular_wall(puts_agg, preco_atual, 'put'),
        'grade_agregada':    grade_agregada,
        'expiracoes_agregadas': sorted(set(f"{r['mes']}{r['ano']}" for r in registros)),
        'fonte':             'Profit/Genial (grade de opções)',
        'vencimento_opcoes': f'{mes_alvo}{ano_alvo}',
        'total_oi_calls':    int(sum(calls.values())),
        'total_oi_puts':     int(sum(puts.values())),
        'arquivo':           os.path.basename(path),
        'grade':             grade,
    }


def analisar_todas_expiracoes(pasta: str) -> dict:
    """
    Ponto de entrada NOVO (11/07/2026) — pedido explícito do usuário: "cada
    contrato tem sua OI... o operador pode selecionar que contrato quer ver
    e o dashboard mostra os dados por contrato". A visão agregada de
    analisar_oi_opcoes() (grade_agregada) mistura todas as expirações num só
    conjunto de paredes — útil para ver os níveis psicológicos do mercado
    como um todo, mas não responde "qual a Call Wall do vencimento X26
    especificamente". Esta função quebra a grade POR expiração individual.

    Para cada expiração presente no arquivo CCM_OP*.xlsx, calcula Call
    Wall/Put Wall/Max Pain usando o ÚLTIMO FECHAMENTO DO PRÓPRIO CONTRATO
    futuro daquele vencimento (CCM<mês><ano>.csv) — não o preço do contrato
    ativo. Motivo: uma parede é "resistência acima do preço atual" — usar o
    preço de um mês diferente classificaria strikes do lado errado (ex.:
    put wall calculado com o preço de U26 aplicado à grade de X26, quando os
    dois têm preços de fechamento diferentes).

    Se não existir arquivo CCM<mês><ano>.csv para aquela expiração de opção
    (grade tem vencimentos sem futuro correspondente listado — pode
    acontecer nas pontas mais longas), usa o strike central da própria
    grade como proxy só para não quebrar o cálculo, e marca
    'preco_estimado': True para o dashboard avisar o operador.
    """
    path = _encontrar_arquivo_op(pasta)
    if path is None:
        return {}

    registros = _extrair_grade(path)
    if not registros:
        return {}

    expiracoes = sorted(set((r['mes'], r['ano']) for r in registros), key=lambda t: _chave_mes_ano(*t))

    saida = {}
    for mes, ano in expiracoes:
        codigo = f'CCM{mes}{ano}'
        grupo = [r for r in registros if r['mes'] == mes and r['ano'] == ano]
        calls = {r['strike']: r['oi'] for r in grupo if r['tipo'] == 'C'}
        puts  = {r['strike']: r['oi'] for r in grupo if r['tipo'] == 'P'}
        todos_strikes = sorted(set(list(calls.keys()) + list(puts.keys())))
        grade = [{'strike': s, 'oi_call': calls.get(s, 0), 'oi_put': puts.get(s, 0)} for s in todos_strikes]

        preco_ref = None
        try:
            preco_ref = ultimo_valor(ler_arquivo(pasta, codigo), 'Close')
        except FileNotFoundError:
            pass

        preco_calc = preco_ref
        preco_estimado = False
        if preco_calc is None and todos_strikes:
            preco_calc = todos_strikes[len(todos_strikes) // 2]
            preco_estimado = True

        saida[codigo] = {
            'contrato':           codigo,
            'vencimento_opcoes':  f'{mes}{ano}',
            'preco_referencia':   round(float(preco_ref), 2) if preco_ref is not None else None,
            'preco_estimado':     preco_estimado,
            'call_wall':          calcular_wall(calls, preco_calc, 'call') if preco_calc is not None else None,
            'put_wall':           calcular_wall(puts, preco_calc, 'put') if preco_calc is not None else None,
            'max_pain':           calcular_max_pain(calls, puts),
            'total_oi_calls':     int(sum(calls.values())),
            'total_oi_puts':      int(sum(puts.values())),
            'grade':              grade,
        }

    return saida


# ── TESTE ─────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    import sys
    pasta = sys.argv[1] if len(sys.argv) > 1 else r'C:\Projetos Phyton\Milho'
    vencimento = sys.argv[2] if len(sys.argv) > 2 else 'CCMU26'
    preco = float(sys.argv[3]) if len(sys.argv) > 3 else 67.25

    print('=' * 60)
    print('TESTE oi_opcoes.py')
    print('=' * 60)
    resultado = analisar_oi_opcoes(pasta, vencimento, preco)
    for k, v in resultado.items():
        print(f'{k}: {v}')

    print()
    print('=' * 60)
    print('TESTE analisar_todas_expiracoes()')
    print('=' * 60)
    todas = analisar_todas_expiracoes(pasta)
    for cod, info in todas.items():
        est = ' (preço estimado)' if info['preco_estimado'] else ''
        print(f"{cod}: preço={info['preco_referencia']}{est} | Call Wall={info['call_wall']} | "
              f"Put Wall={info['put_wall']} | Max Pain={info['max_pain']} | "
              f"OI calls={info['total_oi_calls']} puts={info['total_oi_puts']}")
